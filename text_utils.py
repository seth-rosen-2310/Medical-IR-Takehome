from pdf2image import convert_from_path
import pytesseract
from pytesseract import Output
import numpy as np 
import cv2 
from PIL import Image
import re
import pdfplumber
import pandas as pd
import json
from docx import Document

from llm_utils import *


#checks page for mentions of allergies
def has_allergies(text):
    allergy = re.search(r'(?i)allergies.*', text)
    if allergy is not None:
        return True
    else:
        return False
    
#checks page for mentions of surgeries
def has_surgeries(text):
    surgery = re.search(r'Procedure .* .*\n', text)
    if surgery is not None:
        return True
    else:
        return False
    
#checks page for mentions of medications
def has_medications(text):
    medication = re.search(r'(?i)medication .* started.*\n', text)
    if medication is not None:
        return True
    else:
        return False


#Parses LLM output to turn it into a dictionary for easier use
def parse_allergies(extract_allergies, page_num):
    start = extract_allergies.find("[")
    end = extract_allergies.find("]") + 1

    clean_output = extract_allergies[start : end]
    clean_allergies = json.loads(clean_output)
    for a in clean_allergies:
        a['page'] = page_num

    return clean_allergies

def parse_surgeries(extract_surgeries, page_num):
    start = extract_surgeries.find("[")
    end = extract_surgeries.find("]") + 1

    clean_output = extract_surgeries[start : end]
    clean_surgeries = json.loads(clean_output)
    for s in clean_surgeries:
        s['page'] = page_num

    return clean_surgeries

def parse_medications(extract_medications, page_num):
    start = extract_medications.find("[")
    end = extract_medications.find("]") + 1

    clean_output = extract_medications[start : end]
    clean_medications = json.loads(clean_output)
    for m in clean_medications:
        m['page'] = page_num

    return clean_medications


#combines all of the parsed LLM outputs for finalization
def combine_outputs(parsed_outputs):
    all_outputs = []
    for i in parsed_outputs:
        for j in i:
            all_outputs.append(j)

    return all_outputs


#expect list of dicts from parsed LLM outputs and converts them into dataframes for doc output
def finalize_allergies(allergies):
    return pd.DataFrame.from_dict(allergies)

def finalize_surgeries(surgeries):
    return pd.DataFrame.from_dict(surgeries)
    
def finalize_medications(medications):
    df = pd.DataFrame.from_dict(medications)
    return df.sort_values(by=['drug', 'start date'])


#preprocesses pages that are harder to deal with/lower quality images
def preprocess_image(img):
    grey_scale = cv2.cvtColor(np.array(img),cv2.COLOR_BGR2GRAY)
    resized_image = cv2.resize(grey_scale,None,fx=1.5,fy =1.5,interpolation=cv2.INTER_LINEAR)
    processed_image =cv2.adaptiveThreshold(
    resized_image,255,
    cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
    cv2.THRESH_BINARY,63,13
    ) 
    return processed_image


#extract text from page using one of the OCR libraries
def get_text(pdf, pages, page_num):
    page = pdf.pages[page_num]
    text = page.extract_text(keep_blank_chars=True)
    if text == "":
        img = preprocess_image(pages[page_num])
        text= pytesseract.image_to_string(img,lang='eng')
        
    return text



def add_section(document, df, heading):
    document.add_heading(heading, level=1)

    # add a table to the end and create a reference variable
    # extra row is so we can add the header row
    t = document.add_table(df.shape[0]+1, df.shape[1])

    # add the header rows.
    for j in range(df.shape[-1]):
        t.cell(0,j).text = df.columns[j]

    # add the rest of the data frame
    for i in range(df.shape[0]):
        for j in range(df.shape[-1]):
            t.cell(i+1,j).text = str(df.values[i,j])


def run_extraction(pdf_path):
    #load the pdf to be used by both OCR options depending on page contents
    pages = convert_from_path(pdf_path)
    pdf = pdfplumber.open(pdf_path)

    allergies_temp = []
    surgeries_temp = []
    medications_temp = []


    #loop through whole pdf
    for p in range(len(pdf.pages)):

        text = get_text(pdf, pages, p)

        try:
            if has_allergies(text):
                #extract and parse allergies
                extracted_allergies = gpt_extract_allergies(text)
                parsed_allergies = parse_allergies(extracted_allergies, p)
                allergies_temp.append(parsed_allergies)

            elif has_surgeries(text):
                #extract and parse surgeries
                extracted_surgeries = gpt_extract_surgeries(text)
                parsed_surgeries = parse_surgeries(extracted_surgeries, p)
                surgeries_temp.append(parsed_surgeries)

            elif has_medications(text):
                #extract and parse medications
                extracted_medications = gpt_extract_medications(text)
                parsed_medications = parse_medications(extracted_medications, p)
                medications_temp.append(parsed_medications)

            else:
                continue
        except:
            continue
    
    #combine all parsed LLM outputs
    combined_allergies = combine_outputs(allergies_temp)
    combined_surgeries = combine_outputs(surgeries_temp)
    combined_medications = combine_outputs(medications_temp)

    #finalize outputs in DF for export to doc
    allergies_final = finalize_allergies(combined_allergies)
    surgeries_final = finalize_surgeries(combined_surgeries)
    medications_final = finalize_medications(combined_medications)

    return [allergies_final, surgeries_final, medications_final]


    
